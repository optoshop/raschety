#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Обновить лист Calc!K из shop_*.csv (цены сайта) и перенести старые K → AE.

Правила: docs/pricing-rules.md
"""
from __future__ import annotations

import argparse
import csv
import json
import sys
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import PatternFill
from openpyxl.utils import get_column_letter

QTY_TO_TIER = {
    50: "50-90 шт",
    100: "100-290 шт",
    300: "300-490 шт",
    500: "500-990 шт",
    1000: "1000+ шт",
}

GREEN = PatternFill(fill_type="solid", fgColor="92D050")


def fnum(v):
    if v is None or v == "":
        return None
    try:
        return float(str(v).replace(",", ".").strip())
    except (TypeError, ValueError):
        return None


def money_num(v: float) -> float:
    """Число как в shop money()/toFixed(2) без лишних нулей смысла — float round 2."""
    return round(float(v) + 1e-12, 2)


def norm_qty(q: str) -> str:
    return (q or "").strip().replace("1000 + шт", "1000+ шт")


def load_changed_shop_prices(shop_base: Path, shop_updated: Path) -> dict[tuple[str, str], float]:
    """SKU+tier → новая цена только для parent, у которых regular_price реально менялся."""
    with shop_base.open(newline="", encoding="utf-8") as f:
        base = list(csv.DictReader(f))
    with shop_updated.open(newline="", encoding="utf-8") as f:
        upd = list(csv.DictReader(f))
    if len(base) != len(upd):
        raise SystemExit(f"row count mismatch base={len(base)} upd={len(upd)}")

    changed_parents: set[str] = set()
    for a, b in zip(base, upd):
        parent = (a.get("parent_sku") or "").strip()
        if not parent:
            continue
        if (a.get("regular_price") or "").strip() != (b.get("regular_price") or "").strip():
            changed_parents.add(parent)

    out: dict[tuple[str, str], float] = {}
    for row in upd:
        parent = (row.get("parent_sku") or "").strip()
        if parent not in changed_parents:
            continue
        qty = norm_qty(row.get("meta:attribute_pa_colichestvo") or "")
        if qty not in {
            "50-90 шт",
            "100-290 шт",
            "300-490 шт",
            "500-990 шт",
            "1000+ шт",
        }:
            continue
        n = fnum(row.get("regular_price"))
        if n is None:
            continue
        out[(parent, qty)] = money_num(n)
    return out


def iter_calc_groups(ws, zak_sku: dict[int, str]):
    r = 2
    while r <= ws.max_row:
        a = ws.cell(r, 1).value
        if a is None:
            r += 1
            continue
        try:
            zakup_row = int(a)
        except (TypeError, ValueError):
            r += 1
            continue
        sku = zak_sku.get(zakup_row)
        rows = []
        for i in range(5):
            rr = r + i
            if rr > ws.max_row:
                break
            rows.append(rr)
        yield sku, zakup_row, rows
        r += 5


def process(xlsx_in: Path, shop_base: Path, shop_updated: Path, xlsx_out: Path) -> dict:
    prices = load_changed_shop_prices(shop_base, shop_updated)
    changed_skus = {sku for sku, _ in prices}

    wb = load_workbook(xlsx_in)
    ws = wb["Calc"]
    zak = wb["Закуп"]

    zak_sku: dict[int, str] = {}
    for r in range(2, zak.max_row + 1):
        raw = zak.cell(r, 11).value
        if raw is None or str(raw).strip() == "":
            continue
        zak_sku[r] = str(raw).strip()

    # AE header if empty
    if ws.cell(1, 31).value in (None, ""):
        ws.cell(1, 31).value = "Старая продажа с НДС (до обновления)"

    stats = {
        "k_updated": 0,
        "ae_written": 0,
        "skipped_same": 0,
        "skipped_no_shop_tier": 0,
        "groups_touched": 0,
        "examples": [],
        "skus_touched": [],
    }

    for sku, zakup_row, rows in iter_calc_groups(ws, zak_sku):
        if not sku or sku not in changed_skus:
            continue
        touched = False
        for rr in rows:
            qty_raw = ws.cell(rr, 9).value  # I
            try:
                qty = int(qty_raw)
            except (TypeError, ValueError):
                stats["skipped_no_shop_tier"] += 1
                continue
            tier = QTY_TO_TIER.get(qty)
            if not tier:
                stats["skipped_no_shop_tier"] += 1
                continue
            new_p = prices.get((sku, tier))
            if new_p is None:
                stats["skipped_no_shop_tier"] += 1
                continue
            cell_k = ws.cell(rr, 11)
            old_p = fnum(cell_k.value)
            if old_p is not None and abs(old_p - new_p) < 1e-9:
                stats["skipped_same"] += 1
                continue
            # old → AE only when changed
            ws.cell(rr, 31).value = cell_k.value if cell_k.value is not None else old_p
            stats["ae_written"] += 1
            cell_k.value = new_p
            cell_k.fill = GREEN
            stats["k_updated"] += 1
            touched = True
            if len(stats["examples"]) < 4:
                stats["examples"].append(
                    {
                        "sku": sku,
                        "row": rr,
                        "qty": qty,
                        "tier": tier,
                        "old": old_p,
                        "new": new_p,
                    }
                )
        if touched:
            stats["groups_touched"] += 1
            stats["skus_touched"].append(sku)

    xlsx_out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(xlsx_out)
    stats["skus_touched_count"] = len(stats["skus_touched"])
    stats["out"] = str(xlsx_out)
    return stats


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--xlsx-in", type=Path, default=Path("1. Коробки  Закуп + ОПТ.xlsx"))
    ap.add_argument("--shop-base", type=Path, default=Path("shop_27.10.csv"))
    ap.add_argument("--shop-updated", type=Path, default=Path("shop_27.10-updated.csv"))
    ap.add_argument("--xlsx-out", type=Path, default=Path("1. Коробки  Закуп + ОПТ-updated.xlsx"))
    ap.add_argument("--stats-out", type=Path, default=None)
    args = ap.parse_args(argv)

    stats = process(args.xlsx_in, args.shop_base, args.shop_updated, args.xlsx_out)
    if args.stats_out:
        args.stats_out.write_text(json.dumps(stats, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(stats, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    sys.exit(main())
